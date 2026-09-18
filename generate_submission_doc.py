import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def build_natural_submission_document():
    doc = docx.Document()

    # Document Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Clean, professional color scheme
    COLOR_PRIMARY = RGBColor(26, 77, 46)     # Deep Forest Green #1A4D2E
    COLOR_SECONDARY = RGBColor(30, 41, 59)   # Slate #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500 #64748B

    # Header
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Darukaa.Earth — Full-Stack Hackathon Submission")
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Technical Overview, System Architecture & Setup Guide")
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = COLOR_MUTED
    p_sub.paragraph_format.space_after = Pt(12)

    # Candidate & Submission Information Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Candidate", "Pankaj Sahu"),
        ("Role", "Full-Stack Developer Challenge (Darukaa.Earth)"),
        ("GitHub Repository", "https://github.com/Pankajsahu1234/Darukaa-Project.git"),
        ("Local URLs", "Frontend: http://localhost:5173  |  API Docs: http://localhost:8000/docs"),
        ("Submission Date", "September 18, 2026"),
    ]
    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(5.0)
        c0.text = label
        c1.text = val
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. Repository & Access
    h1 = doc.add_heading("1. Repository & Collaborator Access", level=1)
    h1.paragraph_format.space_before = Pt(12)

    p1 = doc.add_paragraph()
    p1.add_run("GitHub Repository: ").bold = True
    p1.add_run("https://github.com/Pankajsahu1234/Darukaa-Project.git\n\n")
    p1.add_run(
        "The project has been developed with a clean, incremental git commit history across both frontend and backend modules:\n"
        " • Initial project structure, configuration, and CI/CD pipelines\n"
        " • Backend API with PostGIS spatial database, JWT auth, and CRUD endpoints\n"
        " • Frontend application with Mapbox GL JS, polygon drawing tool, and Chart.js analytics\n"
        " • Seed CLI with realistic biomes (Sundarbans, Western Ghats, Amazon, Congo Basin)\n"
        " • Evaluator inspection views and role-based access control\n\n"
        "Collaborator invitations have been added for the review team:\n"
        " • ankita.dasgupta@darukaa.com\n"
        " • harsh.kumar@darukaa.com\n"
        " • utkarsh.gauniyal@darukaa.com\n"
        " • guneet.mutreja@darukaa.com"
    )

    # 2. Review Accounts
    h2 = doc.add_heading("2. Review Accounts & Credentials", level=1)
    h2.paragraph_format.space_before = Pt(12)

    doc.add_paragraph(
        "The database seed script creates three accounts to test different roles. The login page also includes quick auto-fill buttons so you don't have to type credentials manually:"
    )

    cred_table = doc.add_table(rows=1, cols=4)
    cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hdrs = cred_table.rows[0].cells
    c_hdrs[0].text = "Role"
    c_hdrs[1].text = "Email"
    c_hdrs[2].text = "Password"
    c_hdrs[3].text = "Intended Testing Flow"
    for c in c_hdrs:
        set_cell_background(c, "E2E8F0")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(9.0)

    cred_rows = [
        ("Admin", "admin@darukaa.earth", "Password123!", "Full access: create projects, draw polygon sites on the map, add telemetry"),
        ("Evaluator", "evaluator@darukaa.earth", "Password123!", "Audit view: evaluator status banners, verification badges, and scorecard"),
        ("Viewer", "viewer@darukaa.earth", "Password123!", "Read-only inspection of projects, map layers, and time-series charts"),
    ]
    for r1, r2, r3, r4 in cred_rows:
        rc = cred_table.add_row().cells
        rc[0].text = r1
        rc[1].text = r2
        rc[2].text = r3
        rc[3].text = r4
        for cell in rc:
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            set_cell_margins(cell, 60, 60, 80, 80)

    # 3. How to Run Locally
    h3 = doc.add_heading("3. How to Run Locally", level=1)
    h3.paragraph_format.space_before = Pt(12)

    doc.add_paragraph(
        "You can run the application either directly on your host machine or via Docker Compose. Both methods populate initial seed data so you have realistic data ready immediately."
    )

    p_native = doc.add_paragraph()
    p_native.add_run("Method 1: Native Run (Quickest for Local Development)\n").bold = True
    p_native.add_run(
        "Backend (Python 3.10+ / FastAPI):\n"
        "  1. cd backend\n"
        "  2. pip install -r requirements.txt\n"
        "  3. python seed.py              # Creates tables and seeds 4 global biomes with 12 months of data\n"
        "  4. python -m uvicorn app.main:app --host 127.0.0.1 --port 8000 --reload\n"
        "  -> API available at http://localhost:8000 (Swagger docs at http://localhost:8000/docs)\n\n"
        "Frontend (Node 18+ / Vite + React):\n"
        "  1. cd frontend\n"
        "  2. npm install\n"
        "  3. npm run dev\n"
        "  -> Dashboard available at http://localhost:5173\n"
    )

    p_docker = doc.add_paragraph()
    p_docker.add_run("Method 2: Docker Compose (Full PostGIS Container)\n").bold = True
    p_docker.add_run(
        "  1. docker-compose up -d        # Starts PostgreSQL 16 + PostGIS 3.4, Backend, and Frontend\n"
        "  2. docker-compose exec backend python seed.py\n"
        "  -> Open http://localhost:5173 in your browser\n"
    )

    # 4. System Architecture
    h4 = doc.add_heading("4. System Architecture & Tech Stack", level=1)
    h4.paragraph_format.space_before = Pt(12)

    doc.add_paragraph(
        "The application uses a decoupled client-server architecture with a strict separation of concerns between spatial calculation, data querying, and client-side rendering:"
    )

    arch_table = doc.add_table(rows=1, cols=3)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    a_hdrs = arch_table.rows[0].cells
    a_hdrs[0].text = "Component"
    a_hdrs[1].text = "Technology"
    a_hdrs[2].text = "Implementation Role"
    for c in a_hdrs:
        set_cell_background(c, "E2E8F0")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(9.0)

    arch_items = [
        ("Frontend UI", "React 18 + Vite (TypeScript strict)", "Single-page application with type-safe state, component modularity, and fast HMR."),
        ("Map & Spatial Tools", "Mapbox GL JS v3 + @mapbox/mapbox-gl-draw", "Vector map with satellite basemap, custom polygon styling, centroid labels, and interactive polygon creation tool."),
        ("Analytics Charts", "Chart.js (react-chartjs-2)", "Multi-metric time-series visualizer (line curves and monthly volume bars) with 3m, 6m, 12m filters and growth rate metrics."),
        ("Client State & Caching", "TanStack Query v5 + Zustand", "Server-state caching, background re-validation, and optimistic updates handled via TanStack Query; auth and UI draw state managed by Zustand."),
        ("Backend Framework", "FastAPI (Python 3.12, Async)", "Async request handling with SQLAlchemy 2.0 async sessions. Layered architecture: routers -> services -> models -> schemas."),
        ("Spatial Database", "PostgreSQL 16 + PostGIS 3.4", "Stores polygon boundaries in GEOMETRY(Polygon, 4326) with spatial GIST indexes. Computes accurate surface areas using PostGIS ST_Area(geography(geom))."),
        ("Authentication", "JWT (python-jose + passlib/bcrypt)", "Standard Bearer token flow with 60-minute access token lifespan and role-based route protection."),
        ("CI / CD Pipelines", "GitHub Actions + Pre-commit (Husky)", "Automated linting (Ruff, ESLint), formatting (Black), type checking (tsc), and integration tests on push."),
    ]
    for comp, tech, role in arch_items:
        row = arch_table.add_row().cells
        row[0].text = comp
        row[1].text = tech
        row[2].text = role
        for cell in row:
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            set_cell_margins(cell, 50, 50, 70, 70)

    # 5. Database Schema & PostGIS Modeling
    h5 = doc.add_heading("5. Database Schema & Spatial Calculations", level=1)
    h5.paragraph_format.space_before = Pt(12)

    doc.add_paragraph(
        "The relational database schema is designed around projects, geographic site polygons, and time-series telemetry:"
    )

    schema_table = doc.add_table(rows=1, cols=4)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_hdrs = schema_table.rows[0].cells
    s_hdrs[0].text = "Table"
    s_hdrs[1].text = "Columns"
    s_hdrs[2].text = "Indexes & Foreign Keys"
    s_hdrs[3].text = "Description"
    for c in s_hdrs:
        set_cell_background(c, "E2E8F0")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(9.0)

    schema_rows = [
        ("users", "id (UUID), email (string), hashed_password, full_name, role, created_at", "UNIQUE(email), INDEX(id)", "Stores platform accounts and RBAC roles (admin, evaluator, viewer)."),
        ("projects", "id (UUID), name, description, owner_id, created_at, updated_at", "INDEX(owner_id), FK -> users(id) ON DELETE CASCADE", "Top-level environmental projects grouping multiple geographic parcels."),
        ("sites", "id (UUID), project_id, name, geom (Polygon, 4326), area_hectares, site_type, created_at", "GIST(geom), INDEX(project_id), FK -> projects(id) ON DELETE CASCADE", "Specific boundary parcels with PostGIS spatial geometry and calculated surface area."),
        ("analytics_records", "id (UUID), site_id, metric_name, value, recorded_at, metric_metadata (JSONB)", "INDEX(site_id, metric_name, recorded_at), FK -> sites(id) ON DELETE CASCADE", "Time-series observations: carbon stock, biodiversity score, canopy cover, soil carbon."),
    ]
    for t, col, ix, desc in schema_rows:
        row = schema_table.add_row().cells
        row[0].text = t
        row[1].text = col
        row[2].text = ix
        row[3].text = desc
        for cell in row:
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            set_cell_margins(cell, 50, 50, 70, 70)

    p_postgis = doc.add_paragraph()
    p_postgis.paragraph_format.space_before = Pt(8)
    p_postgis.add_run("Why PostGIS ST_Area(geography(geom)) Matters:\n").bold = True
    p_postgis.add_run(
        "Standard planar Euclidean math (such as calculating area on Web Mercator EPSG:3857 coordinates) distorts heavily outside the equator. In higher latitudes, planar calculations can overestimate parcel area by 200% to 400%.\n"
        "To guarantee real-world accuracy for carbon credit verification, Darukaa.Earth calculates area using PostGIS ST_Area(geography(geom)) / 10000.0, which computes true geodesic surface area on the WGS 84 ellipsoid in hectares.\n"
        "For local testing without requiring PostGIS installation, a custom GeoPolygon TypeDecorator cleanly falls back to spherical excess math on SQLite so all test suites run out of the box."
    )

    # 6. CI/CD & Testing
    h6 = doc.add_heading("6. CI/CD Pipeline & Automated Testing", level=1)
    h6.paragraph_format.space_before = Pt(12)

    doc.add_paragraph(
        "Continuous integration is managed through GitHub Actions workflows in .github/workflows/:\n"
        " • backend-ci.yml: Runs on every push and PR. Sets up Python 3.12 with a live PostGIS service container, runs Ruff for linting, Black for code formatting, and executes Pytest with coverage across auth, project creation, site geometries, and analytics.\n"
        " • frontend-ci.yml: Runs on frontend changes. Executes ESLint, strict TypeScript compiler verification (tsc --noEmit), and verifies production Vite bundling.\n"
        " • deploy.yml: Pre-configured deployment triggers for Vercel (frontend) and Render (backend).\n"
        " • Local Pre-commit: Husky and lint-staged ensure that formatting and linting errors are caught before code is committed."
    )

    # 7. Key Engineering Trade-offs & Design Choices
    h7 = doc.add_heading("7. Key Engineering Decisions & Trade-offs", level=1)
    h7.paragraph_format.space_before = Pt(12)

    doc.add_paragraph(
        "Here are the core technical trade-offs evaluated during development:"
    )

    decisions = [
        ("TanStack Query vs. Redux Toolkit",
         "I opted for TanStack Query (React Query) rather than a heavy Redux setup. Most of the platform's state is server-driven (projects, sites, analytics). TanStack Query provides built-in caching, background re-fetching, and cache invalidation without writing hundreds of lines of action-creator and reducer boilerplate. Zustand was added for the tiny amount of client-only state (such as drawing tool toggles and auth tokens)."),

        ("PostGIS Geodesic Area vs. Client-side Turf.js",
         "While turf.area() can compute area on the client for instant UI feedback, client calculations cannot be trusted as an authoritative record for carbon credit issuance. In Darukaa.Earth, the client calculates an instant estimate as the user draws points, but the backend validates the GeoJSON and computes the authoritative surface area in hectares using PostGIS ST_Area(geography(geom)) before saving to the database."),

        ("Chart.js vs. D3.js",
         "D3 is flexible but requires substantial custom code for basic interactions, tooltips, and responsiveness. For multi-metric environmental dashboards, Chart.js (via react-chartjs-2) provides an optimal balance: smooth 60fps canvas rendering, responsive layouts, easily configurable splines, and low bundle overhead."),

        ("FastAPI Async Architecture vs. Django REST Framework",
         "FastAPI was chosen for its native async/await support, Pydantic v2 validation, and automatic OpenAPI schema generation. When dealing with geospatial queries and time-series aggregations, non-blocking I/O allows the server to handle concurrent user requests efficiently with minimal resource usage."),
    ]

    for title, exp in decisions:
        p_dec = doc.add_paragraph()
        p_dec.add_run(f"• {title}: ").bold = True
        p_dec.add_run(exp)
        p_dec.paragraph_format.space_after = Pt(4)

    # 8. Verification Steps for Evaluators
    h8 = doc.add_heading("8. Quick Verification Walkthrough", level=1)
    h8.paragraph_format.space_before = Pt(12)

    doc.add_paragraph(
        "To quickly test the core user stories after starting the app:\n"
        "1. Open http://localhost:5173 and click 'Auto-fill Evaluator' (or 'Auto-fill Admin') and sign in.\n"
        "2. User Story 1 (Project & Site Creation): Click 'New Project' in the navbar. Give it a name and description. Under 'Sites', click 'Draw Polygon on Map' or use a preset biome. Draw 3+ points on the map and double-click to finish. Notice the live calculated hectare size. Click 'Save Project'.\n"
        "3. User Story 2 (Interactive Map): On the main dashboard, all sites render as color-coded polygons with centroid badges. Zoom and pan across biomes (Sundarbans, Western Ghats, Amazon, Congo Basin). Click any polygon or card to open its detail sheet.\n"
        "4. User Story 3 (Site Analytics & Time-Series): When viewing a site's analytics panel, toggle between metrics (Carbon Stock, Biodiversity Index, Canopy Cover, Soil Organic Carbon) and timeframes (3m, 6m, 12m, All). Observe the growth percentage calculation and the monthly breakdown bars."
    )

    # 9. Future Scalability Considerations
    h9 = doc.add_heading("9. Scalability Considerations & Next Steps", level=1)
    h9.paragraph_format.space_before = Pt(12)

    doc.add_paragraph(
        "For taking this platform to enterprise scale with thousands of sites:\n"
        " • Vector Tile Generation (MVT): For 10,000+ polygons, serving raw GeoJSON creates bandwidth bottlenecks. Transitioning to PostGIS ST_AsMVT with tile caching (via Martin or pg_tileserv) would keep map rendering lightning-fast.\n"
        " • Background Task Processing: Heavy calculations (e.g. historical NDVI satellite raster processing) should be offloaded to Celery workers with Redis.\n"
        " • Registry Integrations: Adding automated webhook exports conforming to Verra (VCS) and Gold Standard data schemas for carbon credit certification."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Signoff
    p_end = doc.add_paragraph()
    r_end = p_end.add_run(
        "Thank you for reviewing my submission. Please feel free to reach out if you have any questions during evaluation."
    )
    r_end.font.size = Pt(10)
    r_end.font.italic = True
    r_end.font.color.rgb = COLOR_SECONDARY

    output_filename = "Darukaa_Earth_FullStack_Submission.docx"
    doc.save(output_filename)
    print(f"[OK] Successfully built natural submission document: {output_filename}")


if __name__ == "__main__":
    build_natural_submission_document()
